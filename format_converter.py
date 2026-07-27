# -*- coding: utf-8 -*-
"""多格式测试用例文件转换器：统一转换为 .xlsx。

支持 .xlsm / .xls / .csv / .md / .xmind 五种来源格式，转换产物走现有
openpyxl 解析、编辑、写回流水线。转换尽量保内容（文本、合并单元格），
不承诺保留单元格样式与图片。
"""

import csv
import json
import os
import re
import zipfile
from xml.etree import ElementTree

from openpyxl import Workbook, load_workbook

# 支持转换的扩展名（小写）
CONVERTIBLE_EXT = (".xls", ".xlsm", ".csv", ".md", ".xmind")

# Excel Sheet 名非法字符与长度限制
_SHEET_BAD_CHARS = re.compile(r"[\[\]:*?/\\]")
_SHEET_MAX_LEN = 31


def convert_to_xlsx(src_path, dst_path):
    """把 src_path 转换为 xlsx 写入 dst_path。失败抛 ValueError（含原因）。"""
    ext = os.path.splitext(src_path)[1].lower()
    if ext == ".xlsm":
        _from_xlsm(src_path, dst_path)
    elif ext == ".xls":
        _from_xls(src_path, dst_path)
    elif ext == ".csv":
        _from_csv(src_path, dst_path)
    elif ext == ".md":
        _from_md(src_path, dst_path)
    elif ext == ".xmind":
        _from_xmind(src_path, dst_path)
    else:
        raise ValueError("不支持转换的格式：" + ext)
    return dst_path


def _safe_sheet_name(name, used):
    """清洗 Sheet 名（去非法字符、限长、去重）。"""
    name = _SHEET_BAD_CHARS.sub(" ", str(name or "")).strip() or "Sheet"
    name = name[:_SHEET_MAX_LEN]
    base, idx = name, 2
    while name in used:
        suffix = "(%d)" % idx
        name = base[:_SHEET_MAX_LEN - len(suffix)] + suffix
        idx += 1
    used.add(name)
    return name


def _from_xlsm(src, dst):
    """xlsm：openpyxl 直接读，另存 xlsx（宏自然剥离）。"""
    try:
        wb = load_workbook(src)
    except Exception as exc:
        raise ValueError("读取 xlsm 失败：%s" % exc)
    wb.save(dst)


def _from_xls(src, dst):
    """xls：xlrd 逐格搬运文本与合并单元格。样式/图片不保留。"""
    try:
        import xlrd
    except ImportError:
        raise ValueError("缺少 xlrd 依赖，无法读取 .xls，请先 pip install xlrd")
    try:
        # formatting_info=True 才会填充 merged_cells；个别损坏文件不支持时回退
        try:
            book = xlrd.open_workbook(src, formatting_info=True)
        except Exception:
            book = xlrd.open_workbook(src, formatting_info=False)
    except Exception as exc:
        raise ValueError("读取 xls 失败：%s" % exc)

    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    for sheet in book.sheets():
        ws = wb.create_sheet(_safe_sheet_name(sheet.name, used))
        for r in range(sheet.nrows):
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                value = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        value = xlrd.xldate_as_datetime(value, book.datemode)
                    except Exception:
                        pass
                elif cell.ctype == xlrd.XL_CELL_NUMBER and value == int(value):
                    value = int(value)
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    value = None
                if value is not None and value != "":
                    ws.cell(row=r + 1, column=c + 1).value = value
        # merged_cells 的边界为半开区间 [rlo, rhi)
        for rlo, rhi, clo, chi in sheet.merged_cells:
            ws.merge_cells(start_row=rlo + 1, end_row=rhi,
                           start_column=clo + 1, end_column=chi)
    if not wb.sheetnames:
        raise ValueError("xls 文件中没有工作表")
    wb.save(dst)


def _read_text(src):
    """按 utf-8-sig 读文本，失败回退 gbk。"""
    try:
        with open(src, encoding="utf-8-sig") as fh:
            return fh.read()
    except UnicodeDecodeError:
        with open(src, encoding="gbk", errors="replace") as fh:
            return fh.read()


def _from_csv(src, dst):
    """csv：单 Sheet 逐行写入，编码 utf-8-sig 回退 gbk。"""
    text = _read_text(src)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    rows = list(csv.reader(text.splitlines()))
    if not any(any(v.strip() for v in row) for row in rows):
        raise ValueError("CSV 文件内容为空")
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            if value != "":
                ws.cell(row=r, column=c).value = value
    wb.save(dst)


# Markdown 表格分隔行：| --- | :---: | 之类
_MD_SEP = re.compile(r"^\s*\|?[\s:|-]+\|?\s*$")
_MD_HEADING = re.compile(r"^\s{0,3}#{1,6}\s+(.+?)\s*#*\s*$")


def _md_split_row(line):
    """把 Markdown 表格行拆为单元格列表（去掉首尾管道）。"""
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    # 支持 \| 转义的管道
    cells = re.split(r"(?<!\\)\|", line)
    return [c.replace("\\|", "|").replace("<br>", "\n").strip() for c in cells]


def _from_md(src, dst):
    """md：解析管道表格，多表格多 Sheet（表格前最近标题作 Sheet 名）。"""
    lines = _read_text(src).splitlines()
    tables = []          # [(标题, [行cells])]
    last_heading = ""
    i = 0
    while i < len(lines):
        m = _MD_HEADING.match(lines[i])
        if m:
            last_heading = m.group(1)
            i += 1
            continue
        # 表格起点：本行含管道，下一行是分隔线
        if "|" in lines[i] and i + 1 < len(lines) and _MD_SEP.match(lines[i + 1]) \
                and "|" in lines[i + 1]:
            rows = [_md_split_row(lines[i])]
            i += 2
            while i < len(lines) and "|" in lines[i] and lines[i].strip():
                rows.append(_md_split_row(lines[i]))
                i += 1
            tables.append((last_heading, rows))
            continue
        i += 1
    if not tables:
        raise ValueError("未在 Markdown 中找到表格（仅支持 | 分隔的管道表格）")

    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    for idx, (heading, rows) in enumerate(tables, start=1):
        ws = wb.create_sheet(_safe_sheet_name(heading or ("表格%d" % idx), used))
        for r, row in enumerate(rows, start=1):
            for c, value in enumerate(row, start=1):
                if value != "":
                    ws.cell(row=r, column=c).value = value
    wb.save(dst)


def _from_xmind(src, dst):
    """xmind：一级分支 → Sheet，根到叶路径 → 行，层级依次为列。"""
    try:
        zf = zipfile.ZipFile(src)
    except Exception as exc:
        raise ValueError("读取 xmind 失败（非法 zip）：%s" % exc)
    with zf:
        names = zf.namelist()
        if "content.json" in names:
            roots = _xmind_roots_json(zf.read("content.json"))
        elif "content.xml" in names:
            roots = _xmind_roots_xml(zf.read("content.xml"))
        else:
            raise ValueError("xmind 文件缺少 content.json/content.xml，无法解析")

    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    for root in roots:
        branches = root.get("children") or []
        # 无一级分支时，把根主题本身作为唯一 Sheet
        if not branches:
            branches = [root]
        for branch in branches:
            ws = wb.create_sheet(_safe_sheet_name(branch.get("title"), used))
            paths = []
            _collect_paths(branch.get("children") or [], [], paths)
            if not paths:
                paths = [[branch.get("title") or ""]]
            depth = max(len(p) for p in paths)
            headers = ["层级%d" % n for n in range(1, depth + 1)] + ["测试结果"]
            for c, h in enumerate(headers, start=1):
                ws.cell(row=1, column=c).value = h
            for r, path in enumerate(paths, start=2):
                for c, value in enumerate(path, start=1):
                    if value:
                        ws.cell(row=r, column=c).value = value
    if not wb.sheetnames:
        raise ValueError("xmind 文件中没有主题内容")
    wb.save(dst)


def _collect_paths(nodes, prefix, out):
    """深度优先收集根到叶的标题路径。"""
    for node in nodes:
        path = prefix + [node.get("title") or ""]
        children = node.get("children") or []
        if children:
            _collect_paths(children, path, out)
        else:
            out.append(path)


def _xmind_roots_json(data):
    """XMind Zen/2020+ 的 content.json → 统一 {title, children} 树。"""
    try:
        sheets = json.loads(data.decode("utf-8"))
    except Exception as exc:
        raise ValueError("解析 xmind content.json 失败：%s" % exc)

    def walk(topic):
        children = ((topic.get("children") or {}).get("attached")) or []
        return {"title": topic.get("title") or "",
                "children": [walk(t) for t in children]}

    roots = []
    for sheet in sheets if isinstance(sheets, list) else []:
        root = sheet.get("rootTopic")
        if root:
            roots.append(walk(root))
    if not roots:
        raise ValueError("xmind 文件中没有主题内容")
    return roots


def _xmind_roots_xml(data):
    """XMind 旧版的 content.xml → 统一 {title, children} 树。"""
    # 安全加固：拒绝 DTD/实体声明，防实体膨胀与外部实体攻击
    head = data[:4096].lstrip()
    if b"<!DOCTYPE" in head or b"<!ENTITY" in head:
        raise ValueError("xmind content.xml 含 DTD/实体声明，已拒绝解析")
    try:
        tree = ElementTree.fromstring(data)
    except Exception as exc:
        raise ValueError("解析 xmind content.xml 失败：%s" % exc)

    def local(tag):
        return tag.rsplit("}", 1)[-1]

    def walk(topic):
        title, children = "", []
        for child in topic:
            name = local(child.tag)
            if name == "title":
                title = child.text or ""
            elif name == "children":
                for topics in child:
                    # 仅取 attached 主干分支，忽略 callout/浮动主题
                    if topics.get("type") in (None, "attached"):
                        children.extend(walk(t) for t in topics
                                        if local(t.tag) == "topic")
        return {"title": title, "children": children}

    roots = []
    for sheet in tree:
        if local(sheet.tag) != "sheet":
            continue
        for node in sheet:
            if local(node.tag) == "topic":
                roots.append(walk(node))
    if not roots:
        raise ValueError("xmind 文件中没有主题内容")
    return roots
